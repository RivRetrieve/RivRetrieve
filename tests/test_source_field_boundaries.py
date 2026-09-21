"""Source-boundary regressions mutate exact recordings, never positive source evidence."""

import json
from dataclasses import replace
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

import rivretrieve as rr
import rivretrieve._internal.discovery as discovery
from rivretrieve._internal.recordings import ReplayTransport, read_recording

DATA = Path(__file__).parent / "test_data"


class AlteredResponseTransport:
    def __init__(self, recordings, alter):
        self.replay = ReplayTransport(recordings)
        self.alter = alter

    def send(self, request):
        response = self.replay.send(request)
        return replace(response, content=self.alter(request, response.content))


def extra_sheet(content):
    workbook = openpyxl.load_workbook(BytesIO(content))
    workbook.copy_worksheet(workbook.worksheets[0]).title = "Unestablished additional sheet"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def invalid_workbook_time(content):
    workbook = openpyxl.load_workbook(BytesIO(content))
    workbook.worksheets[0]["A9"] = "not a source timestamp"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_bosnia_invalid_row_cannot_become_successful_coverage(monkeypatch, tmp_path):
    monkeypatch.setenv("RIVRETRIEVE_CACHE_DIR", str(tmp_path / "cache"))
    recordings = tuple(
        read_recording(DATA / name)
        for name in ("ba_fhmzbih_metadata_index.recording.json", "ba_fhmzbih_2101-B_Q_1Y.recording.json")
    )
    transport = AlteredResponseTransport(
        recordings, lambda request, body: invalid_workbook_time(body) if request.url.endswith(".xlsx") else body
    )
    monkeypatch.setattr(discovery, "HttpClient", lambda: transport)
    selection = rr.find(provider="ba_fhmzbih", station="2101-B", quantity="discharge")
    result = rr.fetch(selection, start="2026-09-01", end="2026-09-03", cache="refresh", on_issue="ignore")
    assert result.data.is_empty()
    assert any(outcome.status == "unsupported" for outcome in result.outcomes)
    assert not any(outcome.status in ("success", "empty") for outcome in result.outcomes)


def test_bosnia_does_not_silently_accept_first_of_multiple_sheets(monkeypatch, tmp_path):
    monkeypatch.setenv("RIVRETRIEVE_CACHE_DIR", str(tmp_path / "cache"))
    recordings = tuple(
        read_recording(DATA / name)
        for name in ("ba_fhmzbih_metadata_index.recording.json", "ba_fhmzbih_2101-B_Q_1Y.recording.json")
    )
    transport = AlteredResponseTransport(
        recordings, lambda request, body: extra_sheet(body) if request.url.endswith(".xlsx") else body
    )
    monkeypatch.setattr(discovery, "HttpClient", lambda: transport)
    selection = rr.find(provider="ba_fhmzbih", station="2101-B", quantity="discharge")
    result = rr.fetch(
        selection, start="2026-09-01", end="2026-09-03", cache="refresh", receipts=True, on_issue="ignore"
    )
    assert result.data.is_empty()
    assert any(outcome.status == "unsupported" and "worksheet" in outcome.reason for outcome in result.outcomes)
    assert not any(outcome.status in ("success", "empty") for outcome in result.outcomes)
    assert len(result.receipts.entries) == 2


@pytest.mark.parametrize("bad_index", [b"{", b"{}", b"[]"])
def test_bosnia_external_index_failure_is_retained_not_fatal(monkeypatch, bad_index):
    recording = read_recording(DATA / "ba_fhmzbih_metadata_index.recording.json")
    transport = AlteredResponseTransport((recording,), lambda request, body: bad_index)
    monkeypatch.setattr(discovery, "HttpClient", lambda: transport)
    result = rr.fetch(
        rr.find(provider="ba_fhmzbih", station="2101-B", quantity="discharge"),
        start="2026-09-01",
        end="2026-09-03",
        receipts=True,
        on_issue="ignore",
    )
    assert result.data.is_empty()
    assert any(outcome.status == "unsupported" for outcome in result.outcomes)
    assert result.issues
    assert result.receipts.entries[0].content == bad_index


@pytest.mark.parametrize("next_value", [123, "https://example.org/untrusted"])
def test_france_invalid_pagination_preserves_independent_series(monkeypatch, next_value):
    recordings = tuple(
        read_recording(DATA / name)
        for name in (
            "fr_hubeau_1011000101_QmnJ_padded.recording.json",
            "fr_hubeau_1011000101_QIXnJ_padded.recording.json",
        )
    )

    def alter(request, body):
        if dict(request.params or {}).get("grandeur_hydro_elab") == "QmnJ":
            document = json.loads(body)
            document["next"] = next_value
            return json.dumps(document).encode()
        return body

    transport = AlteredResponseTransport(recordings, alter)
    monkeypatch.setattr(discovery, "HttpClient", lambda: transport)
    selection = rr.find(provider="fr_hubeau", station="1011000101", quantity="discharge", frequency="daily")
    result = rr.fetch(selection, start="2025-01-03", end="2025-01-03", receipts=True, on_issue="ignore")
    assert set(result.data["product_id"]) == {"discharge_daily_max"}
    assert any(
        outcome.status == "unsupported" and outcome.product_id == "discharge_daily_mean" for outcome in result.outcomes
    )
    assert len(result.receipts.entries) == 2


def test_swiss_catalogue_routes_do_not_select_a_preferred_field():
    from rivretrieve._internal.providers.ch_foen.generate_catalogue import build_products

    routes = dict(build_products().select("product_id", "native_id").iter_rows())
    assert routes["discharge_reported"] == "flow,flow_ls"
    assert routes["stage_reported"] == "height_abs,height"


def test_swiss_active_lineage_describes_identified_observations():
    from rivretrieve._internal.providers.ch_foen.origins import build_acquisition_provenance

    document = repr(build_acquisition_provenance())
    assert "five_column" not in document
    assert "five-column" not in document
    assert "observation.source_series_shape" in document


def test_bosnia_catalogue_retains_source_discharge_claim_without_workbook_id_alias():
    selection = rr.find(provider="ba_fhmzbih", station="4024", quantity="discharge")
    claims = [claim for inventory in selection.inventories for claim in inventory.catalogue_claims]
    assert len(claims) == 1
    assert claims[0].identity.namespace == "wiski.L1_ts_id"
    assert claims[0].identity.published_id
    assert claims[0].identity.description == "81 Web Kontinuirani"
    assert selection.series[0].identity.namespace == "ba_fhmzbih/Q"
    assert selection.series[0].identity.published_id == "81 Web Kontinuirani"
    assert claims[0].identity != selection.series[0].identity


@pytest.mark.parametrize("quantity", ["discharge", "stage"])
def test_hydroportail_title_establishes_statistic_not_frequency(quantity):
    selected = rr.find(provider="fr_hydroportail", station="1232000101", quantity=quantity, statistic="instantaneous")
    assert selected.series
    assert all(item.facts[0].statistic.value == "instantaneous" for item in selected.series)
    assert all(item.facts[0].frequency.value is None for item in selected.series)


def test_hydroportail_contradictory_temporal_title_is_unsupported():
    from rivretrieve._internal.providers.fr_hydroportail.config import config
    from rivretrieve._internal.providers.fr_hydroportail.parse import parse
    from tests.test_fr_hydroportail_station import _empty_payload

    payload = _empty_payload()
    document = json.loads(payload.content)
    document["series"]["title"] = "Débit moyen"
    parsed = parse(replace(payload, content=json.dumps(document).encode()), config())
    assert parsed.outcomes[0].status == "unsupported"


def test_temperature_published_unit_spelling_survives_public_result(monkeypatch):
    recordings = tuple(
        read_recording(DATA / f"fr_hubeau_01001336_temp_padded_p{i}.recording.json") for i in range(1, 6)
    )
    monkeypatch.setattr(discovery, "HttpClient", lambda: ReplayTransport(recordings))
    result = rr.fetch(
        rr.find(provider="fr_hubeau", station="01001336", quantity="temperature"),
        start="2008-07-09",
        end="2008-07-10T23:59:59",
    )
    assert result.data["source_unit"].unique().to_list() == ["°C"]


def test_temperature_contradictory_response_unit_is_not_admitted(monkeypatch):
    recordings = tuple(
        read_recording(DATA / f"fr_hubeau_01001336_temp_padded_p{i}.recording.json") for i in range(1, 6)
    )

    def alter(request, body):
        document = json.loads(body)
        for row in document["data"]:
            row["symbole_unite"] = "K"
        return json.dumps(document).encode()

    monkeypatch.setattr(discovery, "HttpClient", lambda: AlteredResponseTransport(recordings, alter))
    result = rr.fetch(
        rr.find(provider="fr_hubeau", station="01001336", quantity="temperature"),
        start="2008-07-09",
        end="2008-07-10T23:59:59",
        on_issue="ignore",
    )
    assert result.data.is_empty()
    assert any(outcome.status == "unsupported" for outcome in result.outcomes)


def test_france_late_bad_continuation_cannot_certify_partial_interval(monkeypatch, tmp_path):
    monkeypatch.setenv("RIVRETRIEVE_CACHE_DIR", str(tmp_path / "cache"))
    recordings = tuple(
        read_recording(DATA / f"fr_hubeau_01001336_temp_padded_p{i}.recording.json") for i in range(1, 6)
    )
    last = recordings[-1].content

    def alter(request, body):
        if body != last:
            return body
        document = json.loads(body)
        document["next"] = 123
        return json.dumps(document).encode()

    calls = []

    class Counted(AlteredResponseTransport):
        def send(self, request):
            calls.append(request.url)
            return super().send(request)

    monkeypatch.setattr(discovery, "HttpClient", lambda: Counted(recordings, alter))
    broad = rr.find(provider="fr_hubeau", station="01001336", quantity="temperature")
    selection = rr.pick(broad, series_id=broad.series[0].series_id)
    kwargs = {"start": "2008-07-09", "end": "2008-07-10T23:59:59", "on_issue": "ignore", "receipts": True}
    partial = rr.fetch(selection, cache="refresh", **kwargs)
    assert not partial.data.is_empty()
    assert any(outcome.status == "unsupported" for outcome in partial.outcomes)
    assert len(partial.receipts.entries) == 5
    assert len(calls) == 5
    reused = rr.fetch(selection, cache="reuse", **kwargs)
    assert len(calls) == 10
    assert any(outcome.status == "unsupported" for outcome in reused.outcomes)


@pytest.mark.parametrize("policy", ["raise", "warn", "ignore"])
def test_bosnia_index_issue_policy_cannot_admit_numbers(monkeypatch, policy):
    from rivretrieve._internal.issues import IssuePolicyError

    recording = read_recording(DATA / "ba_fhmzbih_metadata_index.recording.json")
    monkeypatch.setattr(
        discovery, "HttpClient", lambda: AlteredResponseTransport((recording,), lambda request, body: b"{}")
    )
    selection = rr.find(provider="ba_fhmzbih", station="2101-B", quantity="discharge")
    if policy == "raise":
        with pytest.raises(IssuePolicyError):
            rr.fetch(selection, start="2026-09-01", end="2026-09-03", on_issue=policy)
    elif policy == "warn":
        with pytest.warns(RuntimeWarning):
            result = rr.fetch(selection, start="2026-09-01", end="2026-09-03", on_issue=policy)
        assert result.data.is_empty() and result.issues
    else:
        result = rr.fetch(selection, start="2026-09-01", end="2026-09-03", on_issue=policy)
        assert result.data.is_empty() and result.issues


def test_bosnia_missing_source_route_preserves_independent_station(monkeypatch):
    recordings = tuple(
        read_recording(DATA / name)
        for name in ("ba_fhmzbih_metadata_index.recording.json", "ba_fhmzbih_4024_Q_1Y.recording.json")
    )

    def alter(request, body):
        if not request.url.endswith("index.json"):
            return body
        return json.dumps([row for row in json.loads(body) if row["metadata_station_no"] != "2101-B"]).encode()

    monkeypatch.setattr(discovery, "HttpClient", lambda: AlteredResponseTransport(recordings, alter))
    selection = rr.pick(rr.find(provider="ba_fhmzbih", quantity="discharge"), station=("2101-B", "4024"))
    result = rr.fetch(selection, start="2026-09-01", end="2026-09-02", on_issue="ignore", receipts=True)
    assert set(result.data["station_id"]) == {"4024"}
    bad = [outcome for outcome in result.outcomes if outcome.station_id == "2101-B"]
    assert bad and all(outcome.status == "unsupported" and outcome.series_id is None for outcome in bad)
    assert any(
        b"2101-B" not in receipt.content and receipt.content.startswith(b"[") for receipt in result.receipts.entries
    )


@pytest.mark.parametrize("provider", ["ba_fhmzbih", "fr_hubeau", "fr_hydroportail"])
def test_mapped_products_preserve_exact_optional_physical_facts(provider):
    from importlib import import_module

    from rivretrieve._internal.catalogues.artifact import load_packaged_catalogue_artifact

    declaration = import_module(f"rivretrieve._internal.providers.{provider}.declaration").declaration
    products = load_packaged_catalogue_artifact(declaration.catalogue).products
    mappings = import_module(f"rivretrieve._internal.providers.{provider}.config").SERIES_MAPPINGS
    for row in products.iter_rows(named=True):
        facts = mappings[row["product_id"]].physical_facts()
        assert row["period_type"] == (
            {"instantaneous": "instant"}.get(facts.temporal_support.value, facts.temporal_support.value) or "unknown"
        )
        assert row["period_anchor"] == (facts.timestamp_anchor.value or "unknown")


def test_bosnia_physics_lineage_uses_all_three_published_workbook_headers():
    from pydantic import TypeAdapter

    from rivretrieve._internal.providers.ba_fhmzbih.origins import WorkbookAccessLedger, build_acquisition_provenance

    ledger = TypeAdapter(WorkbookAccessLedger).validate_json(
        Path("maintenance/catalogue/ba_fhmzbih/inventory/baseline_workbook_access.json").read_bytes()
    )
    provenance = build_acquisition_provenance(ledger)
    binding = next(item for item in provenance.fact_bindings if "source.product.native_physics" in item.facts)
    acquisition = next(
        item
        for source in provenance.source_records
        for item in source.acquisitions
        if item.acquisition_id == binding.acquisition_id
    )
    assert len(acquisition.recording_ids) == 3
    assert {url.rsplit("/", 1)[-1] for url in acquisition.requested_from} == {"Q_1Y.xlsx", "H_1Y.xlsx", "Tvode_1Y.xlsx"}

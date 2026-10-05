"""Reconcile GRDC gauge-zero references with the historical Polish inventory."""

from __future__ import annotations

import hashlib
import math
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import polars as pl

from rivretrieve._internal.catalogue_origins import Field
from rivretrieve._internal.catalogues.native import NativeTable
from rivretrieve._internal.catalogues.station_metadata import build_station_metadata as project_metadata
from rivretrieve._internal.issues import FatalContractError


@dataclass(frozen=True, slots=True)
class GaugeZeroElevation:
    """One workbook row, with Excel-decoded metres or the literal ``ND``."""

    station_id: str
    height: float | int | str
    vertical_reference: str


def read_workbook_elevations(path: Path) -> tuple[GaugeZeroElevation, ...]:
    """Read the pinned workbook without exposing its other fields.

    Excel numeric cells use openpyxl's standard integer/float decoding. Formula
    cells are not evaluated. The workbook is later corroboration, not the
    original acquisition of the historical native inventory.
    """
    import openpyxl

    from rivretrieve._internal.providers.pl_imgw.origins import (
        GAUGE_ZERO_HEADER,
        VERTICAL_REFERENCE_HEADER,
        WORKBOOK_BYTE_SIZE,
        WORKBOOK_SHA256,
    )

    body = path.read_bytes()
    if len(body) != WORKBOOK_BYTE_SIZE or hashlib.sha256(body).hexdigest() != WORKBOOK_SHA256:
        raise FatalContractError("Polish workbook byte identity mismatch")
    book = openpyxl.load_workbook(BytesIO(body), read_only=True, data_only=False)
    try:
        if book.sheetnames != ["List1"]:
            raise FatalContractError("Polish workbook sheet identity mismatch")
        sheet = book["List1"]
        headers = [
            (i, row)
            for i, row in enumerate(sheet.iter_rows(max_row=30, values_only=True), 1)
            if GAUGE_ZERO_HEADER in row
        ]
        if len(headers) != 1:
            raise FatalContractError("Polish workbook height header is not unique")
        number, header = headers[0]
        names = ("Station code", GAUGE_ZERO_HEADER, VERTICAL_REFERENCE_HEADER)
        if any(header.count(name) != 1 for name in names):
            raise FatalContractError("Polish workbook required headers mismatch")
        identity_col, height_col, reference_col = (header.index(name) for name in names)
        rows = []
        for row in sheet.iter_rows(min_row=number + 1, values_only=True):
            if all(value is None for value in row):
                continue
            identity = row[identity_col]
            if type(identity) is int:
                identity = str(identity)
            rows.append(GaugeZeroElevation(identity, row[height_col], row[reference_col]))
        return tuple(rows)
    finally:
        book.close()


def build_station_metadata(
    native_table: NativeTable,
    stations: pl.DataFrame,
    elevations: tuple[GaugeZeroElevation, ...],
) -> pl.DataFrame:
    """Attach each workbook reference without rewriting native height strings.

    Numerical reconciliation uses exact equality between the decoded Excel
    number and ``float(native_height)``. No tolerance, rounding or datum shift
    applies. ``ND`` must occur in both height and reference. Invalid or duplicate
    identities, different station coverage and mismatched heights raise
    ``FatalContractError`` without including controlled values in the error.
    """
    from rivretrieve._internal.providers.pl_imgw.origins import (
        STATION_CATALOGUE_ORIGINS,
        STATION_METADATA_FIELDS,
        VERTICAL_REFERENCE_HEADER,
    )

    references = {}
    heights = {}
    for row in elevations:
        identity = row.station_id
        if (
            not isinstance(identity, str)
            or len(identity) != 9
            or not identity.isascii()
            or not identity.isdigit()
            or identity in references
        ):
            raise FatalContractError("Polish workbook station identity invalid or duplicated")
        if row.vertical_reference not in ("EVRF2007", "Kronsztadt", "ND"):
            raise FatalContractError("Polish workbook vertical reference invalid")
        if row.height == "ND":
            if row.vertical_reference != "ND":
                raise FatalContractError("Polish workbook placeholders are not paired")
        elif (
            isinstance(row.height, bool)
            or not isinstance(row.height, (int, float))
            or not math.isfinite(row.height)
            or row.vertical_reference == "ND"
        ):
            raise FatalContractError("Polish workbook height invalid or placeholder unpaired")
        references[identity] = row.vertical_reference
        heights[identity] = row.height
    native = native_table.data
    if (
        native.schema.get("gauge_altitude") != pl.String
        or native.schema.get("gauge_id") != pl.String
        or native["gauge_id"].n_unique() != native.height
        or set(native["gauge_id"].to_list()) != set(references)
        or not references
    ):
        raise FatalContractError("Polish workbook and native station coverage or scalar types differ")
    for identity, value in native.select("gauge_id", "gauge_altitude").iter_rows():
        expected = heights[identity]
        if expected == "ND":
            matches = value == "ND"
        else:
            try:
                matches = value is not None and float(value) == expected
            except (ValueError, TypeError):
                matches = False
        if not matches:
            raise FatalContractError("Polish workbook and native height mismatch")
    enriched = NativeTable(
        native.with_columns(
            pl.Series(VERTICAL_REFERENCE_HEADER, [references[i] for i in native["gauge_id"]], dtype=pl.String)
        )
    )
    station_origin = STATION_CATALOGUE_ORIGINS["station_id"]
    if not isinstance(station_origin, Field):
        raise FatalContractError("Polish station identity requires a native field origin")
    return project_metadata("pl_imgw", enriched, stations, station_origin, STATION_METADATA_FIELDS)

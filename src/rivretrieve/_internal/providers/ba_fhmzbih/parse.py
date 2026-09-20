"""ba_fhmzbih native observations and source-series evidence.

Contributed by: Thiago von Däniken
"""

from datetime import datetime
from io import BytesIO
from math import isfinite

import openpyxl
import polars as pl

from rivretrieve._internal.engine import Payload, ProviderConfig, Rows, RowsSchema, WithIssues
from rivretrieve._internal.issues import FatalContractError, Issue
from rivretrieve._internal.primitives import ProviderId
from rivretrieve._internal.provider_series import NATIVE_SCHEMA, UnsupportedSourceStructureError, parse_mapped_series
from rivretrieve._internal.providers.ba_fhmzbih.config import SERIES_MAPPINGS, BaFhmzbihSourceCoordinates
from rivretrieve._internal.providers.ba_fhmzbih.fetch import BaFhmzbihMetadataCoordinates
from rivretrieve._internal.source_series import ParsedSeries


def _parse_native(payload: Payload, provider_config: ProviderConfig) -> WithIssues[Rows]:
    if isinstance(payload.source_coordinates.value, BaFhmzbihMetadataCoordinates):
        return WithIssues(_empty())
    if len(payload.station_products) != 1:
        raise FatalContractError("ba_fhmzbih workbook payload must contain one station-product pair")
    station, product = payload.station_products[0]
    coordinates = payload.source_coordinates.value
    if not isinstance(coordinates, BaFhmzbihSourceCoordinates):
        raise FatalContractError("ba_fhmzbih workbook has invalid source coordinates")
    try:
        workbook = openpyxl.load_workbook(BytesIO(payload.content), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        headers = tuple(
            tuple(cell for cell in row[:2]) for row in sheet.iter_rows(min_row=1, max_row=8, values_only=True)
        )
        expected = (
            ("#Station Name", headers[0][1]),
            ("#Station Number", station),
            ("#Station Parameter Name", coordinates.source_parameter),
            ("#Timeseries Name", "81 Web Kontinuirani"),
            ("#Unit Symbol", coordinates.source_unit),
            ("#Rows", headers[5][1]),
            ("#", None),
            ("#Timestamp", "Value"),
        )
        if headers != expected or (
            isinstance(headers[5][1], bool)
            or not isinstance(headers[5][1], int | float)
            or not float(headers[5][1]).is_integer()
        ):
            raise UnsupportedSourceStructureError("ba_fhmzbih workbook headers differ from the source contract")
        rows = []
        invalid_count = 0
        for timestamp, value, *_ in sheet.iter_rows(min_row=9, values_only=True):
            if not isinstance(timestamp, datetime):
                invalid_count += 1
                continue
            if isinstance(value, bool):
                raise UnsupportedSourceStructureError("ba_fhmzbih workbook observation value cannot be boolean")
            try:
                native_value = None if value is None else float(value)
            except OverflowError as error:
                raise UnsupportedSourceStructureError(
                    "ba_fhmzbih workbook observation value is not representable as a finite number"
                ) from error
            except (TypeError, ValueError):
                invalid_count += 1
                continue
            if native_value is not None and not isfinite(native_value):
                raise UnsupportedSourceStructureError("ba_fhmzbih workbook observation value must be finite")
            rows.append(
                {
                    "station_id": station,
                    "product_id": product,
                    "time": timestamp,
                    "value": native_value,
                    "time_zone": "unknown",
                }
            )
        workbook.close()
    except (FatalContractError, UnsupportedSourceStructureError):
        raise
    except Exception as error:
        raise UnsupportedSourceStructureError("ba_fhmzbih payload is not a readable workbook") from error
    if len(rows) + invalid_count != headers[5][1]:
        raise UnsupportedSourceStructureError("ba_fhmzbih workbook declared row count differs from its data rows")
    frame = pl.DataFrame(rows, schema=NATIVE_SCHEMA).sort("time")
    issues = []
    if invalid_count:
        issues.append(
            _issue(
                "invalid_workbook_row",
                "Workbook rows with invalid timestamp or value were dropped",
                station,
                invalid_count,
            )
        )
    if frame.is_empty():
        issues.append(_issue("missing_data", "Workbook contains no observation rows", station, 0))
    return WithIssues(frame, tuple(issues))


def _empty() -> Rows:
    return pl.DataFrame(schema=NATIVE_SCHEMA)


def _issue(code: str, message: str, station: str, count: int) -> Issue:
    return Issue(
        severity="warning",
        code=code,
        message=message,
        details={"station_id": station, "row_count": count},
        provider_id=ProviderId("ba_fhmzbih"),
    )


def parse(payload: Payload, provider_config: ProviderConfig) -> ParsedSeries:
    if isinstance(payload.source_coordinates.value, BaFhmzbihMetadataCoordinates):
        return ParsedSeries(pl.DataFrame(schema=RowsSchema.polars_schema), (), (), ())
    return parse_mapped_series(
        payload, provider_config, provider="ba_fhmzbih", mappings=SERIES_MAPPINGS, native_parse=_parse_native
    )
